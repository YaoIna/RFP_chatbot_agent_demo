from collections.abc import Mapping
from datetime import UTC, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


class GraphDocumentRenderer:
    """Render an approved LangGraph draft to a stable, professional DOCX package."""

    _PACKAGE_TIME = (1980, 1, 1, 0, 0, 0)
    _PROPERTY_TIME = datetime(2000, 1, 1, tzinfo=UTC)

    def __init__(self, output_dir: Path) -> None:
        self.output_dir = output_dir

    def render(self, state: Mapping[str, Any]) -> Path:
        if not state.get("approved"):
            raise ValueError("Document generation requires explicit user approval.")
        case_id = str(state["case_id"])
        if (
            not case_id
            or Path(case_id).name != case_id
            or any(not (character.isalnum() or character in "_-") for character in case_id)
        ):
            raise ValueError("The case ID is not safe for an output filename.")
        version = int(state["draft_version"])
        draft = state["draft"]
        sections = draft["sections"]

        document = Document()
        page = document.sections[0]
        page.orientation = WD_ORIENT.PORTRAIT
        page.page_width = Inches(8.5)
        page.page_height = Inches(11)
        page.top_margin = Inches(0.7)
        page.bottom_margin = Inches(0.7)
        page.left_margin = Inches(0.8)
        page.right_margin = Inches(0.8)

        normal = document.styles["Normal"]
        normal.font.name = "Aptos"
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor(31, 41, 55)
        heading = document.styles["Heading 1"]
        heading.font.name = "Aptos Display"
        heading.font.size = Pt(16)
        heading.font.bold = True
        heading.font.color.rgb = RGBColor(0, 0, 0)
        title_style = document.styles["Title"]
        title_style.font.name = "Aptos Display"
        title_style.font.size = Pt(28)
        title_style.font.bold = True
        title_style.font.color.rgb = RGBColor(0, 0, 0)
        title_style_properties = title_style.element.get_or_add_pPr()
        for border in title_style_properties.findall(qn("w:pBdr")):
            title_style_properties.remove(border)

        title = document.add_paragraph("IT Request for Quotation", style="Title")
        title_properties = title._p.get_or_add_pPr()
        for border in title_properties.findall(qn("w:pBdr")):
            title_properties.remove(border)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = document.add_paragraph(
            f"{str(state['confirmed_type']).replace('_', ' ').title()}  |  "
            f"Reference {case_id}  |  Draft {version}"
        )
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph(
            "Suppliers should respond to every applicable requirement and identify any "
            "assumptions, exclusions, or dependencies in their quotation."
        )
        for section_data in sections:
            section_heading = document.add_heading(str(section_data["heading"]), level=1)
            section_heading.paragraph_format.keep_with_next = True
            lines = [
                line for line in str(section_data["body_markdown"]).splitlines() if line.strip()
            ]
            rendered_lines = []
            for line in lines:
                if line.startswith("- "):
                    paragraph = document.add_paragraph(line[2:], style="List Bullet")
                else:
                    paragraph = document.add_paragraph(line.strip())
                paragraph.paragraph_format.keep_together = True
                rendered_lines.append((line, paragraph))
            for index, (line, _) in enumerate(rendered_lines):
                if not line.startswith("- "):
                    continue
                next_is_bullet = index + 1 < len(rendered_lines) and rendered_lines[index + 1][
                    0
                ].startswith("- ")
                if next_is_bullet:
                    continue
                previous_index = index - 1
                if previous_index < 0:
                    continue
                _, previous_paragraph = rendered_lines[previous_index]
                previous_paragraph.paragraph_format.keep_with_next = True

        document.core_properties.title = "IT Request for Quotation"
        document.core_properties.author = "RFQ Agent"
        document.core_properties.identifier = f"{case_id}:v{version}"
        document.core_properties.subject = str(state["confirmed_type"])
        document.core_properties.created = self._PROPERTY_TIME
        document.core_properties.modified = self._PROPERTY_TIME

        self.output_dir.mkdir(parents=True, exist_ok=True)
        output = self.output_dir / f"{case_id}-v{version}.docx"
        with NamedTemporaryFile(suffix=".docx", dir=self.output_dir, delete=False) as handle:
            raw = Path(handle.name)
        with NamedTemporaryFile(suffix=".docx", dir=self.output_dir, delete=False) as handle:
            stable = Path(handle.name)
        try:
            document.save(str(raw))
            self._stabilize_package(raw, stable)
            stable.replace(output)
        finally:
            raw.unlink(missing_ok=True)
            stable.unlink(missing_ok=True)
        return output

    @classmethod
    def _stabilize_package(cls, source: Path, target: Path) -> None:
        with (
            ZipFile(source) as source_zip,
            ZipFile(target, "w", compression=ZIP_DEFLATED, compresslevel=9) as target_zip,
        ):
            for name in sorted(source_zip.namelist()):
                original = source_zip.getinfo(name)
                info = ZipInfo(name, cls._PACKAGE_TIME)
                info.compress_type = ZIP_DEFLATED
                info.external_attr = original.external_attr
                info.create_system = original.create_system
                target_zip.writestr(info, source_zip.read(name))
